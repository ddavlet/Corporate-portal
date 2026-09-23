"""Excel renderers: sheets, header, number formats per units, outline groups, print setup, operations table, names."""
import os
import subprocess
import sys
from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch

from django.conf import settings
from django.test import SimpleTestCase, override_settings
from openpyxl import load_workbook

from apps.modules.reports.layouts import ProfessionalCashflowLayout, ProfessionalPnlLayout
from apps.modules.reports.periods import PeriodSpec, build_column_set
from apps.modules.reports.test_fixtures import OPENING, START_MONTH, TODAY, sample_entries
from apps.modules.reports.statements import build_statement, statement_to_dict
from apps.modules.reports.xlsx_export import (
    BAD,
    GOOD,
    OPERATION_HEADERS,
    PCT_DELTA_FORMAT,
    RATIO_FORMAT,
    UNIT_FORMATS,
    UNIT_LABELS,
    UNITS,
    LinesExport,
    LinesXlsxRenderer,
    StatementExport,
    StatementXlsxRenderer,
    lines_filename,
    statement_filename,
)

YTD = PeriodSpec(kind="ytd", compare="yoy")
SECTION_LABELS = {"rev": "Выручка", "opex": "Операционные расходы"}
OPERATIONS = [
    {
        "entry_id": "revenue:bank:2:0", "date": "2026-08-05", "amount": "1500.00", "section": "revenue",
        "source": "bank", "category": "Поступление в банк", "title": "CLICK: перечисление", "counterparty": "",
        "request_id": None, "channel": "CLICK", "line_id": "rev.bank", "line_label": "Банк", "amortization": None,
    },
    {
        "entry_id": "operational:request:11:2", "date": "2026-08-01", "amount": "100.00", "section": "operational",
        "source": "request", "category": "Маркетинг", "title": "Выставка", "counterparty": "ООО Поставщик",
        "request_id": 11, "channel": "", "line_id": "opex.a1b2c3d4", "line_label": "Маркетинг",
        "amortization": {"index": 2, "count": 3},
    },
]


def statement_data(layout, spec, period_label):
    """The dict `build_statement_for_tenant` returns, built from the shared fixture."""
    column_set = build_column_set(spec, today=TODAY, start_month=START_MONTH)
    statement = build_statement(
        entries=sample_entries(), layout=layout, column_set=column_set, opening_balance=OPENING, start_month=START_MONTH
    )
    data = statement_to_dict(statement)
    data.update(
        template="professional",
        meta={
            "company": "Демо Трейд", "source": "backend", "generated_at": "2026-09-23T14:32:00+05:00",
            "start_month": START_MONTH, "today": TODAY.isoformat(), "period_label": period_label,
        },
        methodology=[{"label": "Период отчёта", "text": "с января 2025 года"}],
        warnings=[],
    )
    return data


def render_statement(data, *, units="m", operations=OPERATIONS):
    export = StatementExport(
        statement=data, operations=operations, section_labels=SECTION_LABELS, units=units, author="Бухгалтер"
    )
    return load_workbook(BytesIO(StatementXlsxRenderer().render(export)))


def render_lines(items, *, query=""):
    export = LinesExport(
        report="pnl", title="Выручка", date_from=date(2026, 8, 1), date_to=date(2026, 8, 31), query=query,
        items=items, total="1600.00", section_labels=SECTION_LABELS, company="Демо Трейд", author="Бухгалтер",
        generated_at=datetime(2026, 9, 23, 14, 32),
    )
    return load_workbook(BytesIO(LinesXlsxRenderer().render(export)))


def row_of(sheet, label):
    for (cell,) in sheet.iter_rows(min_col=1, max_col=1):
        if cell.value == label:
            return cell.row
    raise AssertionError(f"no row labelled {label!r}")


def column_of(sheet, header):
    for cell in sheet[8]:
        if cell.value is not None and str(cell.value).split("\n")[0] == header:
            return cell.column
    raise AssertionError(f"no column {header!r}")


def cell_at(sheet, label, header):
    return sheet.cell(row=row_of(sheet, label), column=column_of(sheet, header))


def column_a(sheet):
    return [cell.value for (cell,) in sheet.iter_rows(min_col=1, max_col=1) if cell.value]


@override_settings(TIME_ZONE="Asia/Tashkent")
class StatementXlsxRendererTests(SimpleTestCase):
    def setUp(self):
        self.workbook = render_statement(statement_data(ProfessionalPnlLayout(), YTD, "янв – 23 сен 2026"))
        self.sheet = self.workbook["Отчёт"]

    def test_workbook_has_report_operations_and_parameters_sheets(self):
        self.assertEqual(self.workbook.sheetnames, ["Отчёт", "Операции", "Параметры"])

    def test_header_names_company_period_comparison_units_and_author(self):
        self.assertEqual(
            [self.sheet[f"A{number}"].value for number in range(1, 7)],
            [
                "Демо Трейд",
                "Отчёт о прибылях и убытках",
                "Период: янв – 23 сен 2026 (01.01.2026 – 23.09.2026)",
                "Сравнение: с тем же периодом прошлого года (01.01.2025 – 23.09.2025)",
                "Единицы: млн сум",
                "Выгружено: 23.09.2026 14:32 · Бухгалтер",
            ],
        )

    def test_amounts_are_exact_and_units_are_only_a_number_format(self):
        for units in UNITS:
            sheet = render_statement(statement_data(ProfessionalPnlLayout(), YTD, "янв – 23 сен 2026"), units=units)["Отчёт"]
            cell = cell_at(sheet, "Итого выручка", "Итого")
            self.assertEqual((cell.value, cell.number_format), (3500, UNIT_FORMATS[units]))
            self.assertEqual(sheet["A5"].value, f"Единицы: {UNIT_LABELS[units]}")

    def test_cells_hold_values_not_formulas(self):
        for sheet in self.workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", f"{sheet.title}!{cell.coordinate}")

    def test_groups_are_outlines_with_totals_below(self):
        levels = {
            label: self.sheet.row_dimensions[row_of(self.sheet, label)].outline_level
            for label in ("Выручка", "Банк", "Касса", "Продажа", "Итого касса", "Итого выручка")
        }
        self.assertEqual(
            levels, {"Выручка": 1, "Банк": 1, "Касса": 2, "Продажа": 2, "Итого касса": 1, "Итого выручка": 0}
        )
        self.assertTrue(self.sheet.sheet_properties.outlinePr.summaryBelow)
        self.assertLess(row_of(self.sheet, "Продажа"), row_of(self.sheet, "Итого выручка"))
        self.assertFalse(self.sheet.row_dimensions[row_of(self.sheet, "Банк")].hidden)

    def test_totals_results_and_margins_are_styled(self):
        total = cell_at(self.sheet, "Итого выручка", "Итого")
        self.assertTrue(total.font.b)
        self.assertEqual(total.border.top.style, "thin")
        ebit = self.sheet.cell(row=row_of(self.sheet, "EBIT (операционная прибыль)"), column=1)
        self.assertTrue(ebit.fill.fgColor.rgb.endswith("EAF1FB"))
        self.assertEqual(cell_at(self.sheet, "Чистая прибыль", "Итого").border.bottom.style, "double")
        margin = cell_at(self.sheet, "маржа EBIT", "Итого")
        self.assertEqual((margin.value, margin.number_format, margin.font.i), (0.8571, RATIO_FORMAT, True))

    def test_changes_are_green_or_red_by_what_they_mean(self):
        revenue = cell_at(self.sheet, "Итого выручка", "Δ")
        expenses = cell_at(self.sheet, "Итого операционные расходы", "Δ")
        self.assertAlmostEqual(revenue.value, 3.375, places=4)
        self.assertEqual(revenue.number_format, PCT_DELTA_FORMAT)
        self.assertTrue(revenue.font.color.rgb.endswith(GOOD))
        self.assertTrue(expenses.font.color.rgb.endswith(BAD))

    def test_unfinished_month_is_explained_below_the_table(self):
        self.assertIn("* Сен: период не закрыт, данные по 23.09.2026.", column_a(self.sheet))

    def test_prints_landscape_one_page_wide_with_repeated_header(self):
        setup = self.sheet.page_setup
        self.assertEqual((setup.orientation, setup.fitToWidth, setup.fitToHeight), ("landscape", 1, 0))
        self.assertTrue(self.sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(self.sheet.print_title_rows, "$8:$8")
        self.assertEqual(self.sheet.freeze_panes, "B9")
        self.assertIn("&P", self.sheet.oddFooter.center.text)

    def test_operations_sheet_is_a_filterable_table(self):
        sheet = self.workbook["Операции"]
        self.assertEqual([cell.value for cell in sheet[1]], list(OPERATION_HEADERS))
        self.assertEqual(
            [cell.value for cell in sheet[3]],
            [datetime(2026, 8, 1), "Операционные расходы", "Маркетинг", "Заявка", 11, "ООО Поставщик", "Выставка",
             100, "2026-08", "2 из 3"],
        )
        self.assertEqual((sheet["B2"].value, sheet["F2"].value), ("Выручка", None))
        self.assertEqual(sheet.tables["Operations"].ref, "A1:J3")
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_parameters_sheet_lists_the_methodology(self):
        sheet = self.workbook["Параметры"]
        self.assertEqual((sheet["A2"].value, sheet["B2"].value), ("Период отчёта", "с января 2025 года"))

    def test_month_pack_names_both_comparisons(self):
        data = statement_data(ProfessionalPnlLayout(), PeriodSpec(kind="month", month="2026-08"), "авг 2026")
        sheet = render_statement(data)["Отчёт"]
        self.assertEqual(sheet["A3"].value, "Период: авг 2026 (01.08.2026 – 31.08.2026)")
        self.assertEqual(sheet["A4"].value, "Сравнение: с предыдущим месяцем и тем же месяцем прошлого года")
        for header in ("Август 2026", "Июль 2026", "Август 2025", "С начала года"):
            column_of(sheet, header)

    def test_periods_before_accounting_are_empty_and_explained(self):
        data = statement_data(ProfessionalPnlLayout(), PeriodSpec(kind="year", year=2025, compare="yoy"), "2025")
        sheet = render_statement(data)["Отчёт"]
        self.assertIsNone(cell_at(sheet, "Итого выручка", "Год назад").value)
        self.assertIn("Пустые ячейки — период до начала учёта, данных нет.", column_a(sheet))

    def test_cash_flow_closing_balance_is_emphasised(self):
        sheet = render_statement(statement_data(ProfessionalCashflowLayout(), YTD, "янв – 23 сен 2026"))["Отчёт"]
        self.assertEqual(sheet["A2"].value, "Отчёт о движении денежных средств")
        closing = cell_at(sheet, "Остаток на конец периода", "Итого")
        self.assertTrue(closing.font.b)
        self.assertEqual(closing.border.bottom.style, "double")


class LinesXlsxRendererTests(SimpleTestCase):
    def test_header_says_what_the_list_is_and_totals_it(self):
        workbook = render_lines(OPERATIONS, query="click")
        self.assertEqual(workbook.sheetnames, ["Операции"])
        sheet = workbook["Операции"]
        self.assertEqual(
            [sheet[f"A{number}"].value for number in range(1, 7)],
            [
                "Демо Трейд",
                "Отчёт о прибылях и убытках · Выручка",
                "Период: 01.08.2026 – 31.08.2026",
                "Поиск: «click»",
                "Итого, сум",
                "Выгружено: 23.09.2026 14:32 · Бухгалтер",
            ],
        )
        self.assertEqual(sheet["B5"].value, 1600)
        self.assertEqual([cell.value for cell in sheet[8]], list(OPERATION_HEADERS))
        self.assertEqual(sheet.tables["Operations"].ref, "A8:J10")
        self.assertEqual(sheet.freeze_panes, "A9")

    def test_writes_a_valid_sheet_when_there_are_no_operations(self):
        sheet = render_lines([])["Операции"]
        self.assertEqual(sheet["A4"].value, "Поиск: —")
        self.assertEqual(sheet["A9"].value, "Операций за период нет")
        self.assertEqual(len(sheet.tables), 0)


class ExportFilenameTests(SimpleTestCase):
    def test_statement_file_is_named_after_report_period_tenant_and_day(self):
        month = PeriodSpec(kind="month", month="2026-08")
        self.assertEqual(statement_filename("pnl", month, date(2026, 8, 31), "demo", TODAY), "PnL_2026-08_demo_2026-09-23.xlsx")
        self.assertEqual(statement_filename("pnl", YTD, date(2026, 9, 23), "demo", TODAY), "PnL_2026-YTD_demo_2026-09-23.xlsx")
        self.assertEqual(
            statement_filename("cashflow", PeriodSpec(kind="year", year=2025), date(2025, 12, 31), "demo", TODAY),
            "Cashflow_2025_demo_2026-09-23.xlsx",
        )
        self.assertEqual(
            statement_filename("pnl", PeriodSpec(kind="ltm"), date(2026, 8, 31), "demo", TODAY),
            "PnL_LTM-2026-08_demo_2026-09-23.xlsx",
        )

    def test_lines_file_names_the_range(self):
        self.assertEqual(
            lines_filename("pnl", date(2026, 8, 1), date(2026, 8, 31), "demo", TODAY),
            "PnL_operations_2026-08-01_2026-08-31_demo_2026-09-23.xlsx",
        )


HOSTILE = dict(
    OPERATIONS[0],
    entry_id="revenue:bank:9:0",
    title='=HYPERLINK("http://example.com","x")',
    counterparty="#N/A",
    line_label="Банк\x0b",
)


@override_settings(TIME_ZONE="Asia/Tashkent")
class UserTextTests(SimpleTestCase):
    """Payer and employee text stays text: never a live formula or error value, never a crashed export."""

    def test_statement_workbook_keeps_user_text_as_text(self):
        data = statement_data(ProfessionalPnlLayout(), YTD, "янв – 23 сен 2026")
        data["methodology"] = [{"label": "Назначения", "text": "=1+1"}]
        workbook = render_statement(data, operations=[HOSTILE])
        operations = workbook["Операции"]
        self.assertEqual((operations["G2"].value, operations["G2"].data_type), ('=HYPERLINK("http://example.com","x")', "s"))
        self.assertEqual((operations["F2"].value, operations["F2"].data_type), ("#N/A", "s"))
        self.assertEqual(operations["C2"].value, "Банк")
        parameters = workbook["Параметры"]
        self.assertEqual((parameters["B2"].value, parameters["B2"].data_type), ("=1+1", "s"))

    def test_lines_workbook_survives_control_characters(self):
        sheet = render_lines([HOSTILE])["Операции"]
        self.assertEqual(sheet["G9"].data_type, "s")
        self.assertEqual(sheet["C9"].value, "Банк")


@override_settings(TIME_ZONE="Asia/Tashkent")
class WorkbookSafetyTests(SimpleTestCase):
    def test_the_author_is_the_workbook_creator(self):
        statement = render_statement(statement_data(ProfessionalPnlLayout(), YTD, "янв – 23 сен 2026"))
        self.assertEqual((statement.properties.creator, render_lines(OPERATIONS).properties.creator), ("Бухгалтер", "Бухгалтер"))

    def test_very_long_lists_are_capped_with_a_note(self):
        with patch("apps.modules.reports.xlsx_export.MAX_OPERATION_ROWS", 1):
            sheet = render_lines(OPERATIONS)["Операции"]
        self.assertEqual(sheet.tables["Operations"].ref, "A8:J9")
        self.assertEqual(sheet["A11"].value, "В файле первые 1 из 2 операций. Сузьте период, чтобы выгрузить все.")
        self.assertEqual(sheet["B5"].value, 1600)  # the header total still covers every operation

    def test_reports_api_loads_without_openpyxl(self):
        """Excel is optional: without openpyxl the statement, lines and legacy report endpoints still import."""
        script = (
            "import sys; sys.modules['openpyxl'] = None; import django; django.setup(); "
            "import apps.modules.reports.urls, apps.modules.reports.views, apps.modules.reports.services; "
            "print('reports-ok')"
        )
        env = {**os.environ, "DJANGO_SETTINGS_MODULE": settings.SETTINGS_MODULE}
        result = subprocess.run(
            [sys.executable, "-c", script], cwd=settings.BASE_DIR, env=env, capture_output=True, text=True, timeout=120
        )
        self.assertEqual(result.returncode, 0, result.stderr[-3000:])
        self.assertIn("reports-ok", result.stdout)
