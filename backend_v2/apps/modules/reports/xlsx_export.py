"""Excel export for statements and operation lists.

Cells hold values, never formulas: Telegram and phone previews show the file without recalculating it.
Amounts stay exact сум; the chosen units are only a number format.
"""
from __future__ import annotations

from abc import ABC, abstractmethod
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import Outline, PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from apps.modules.reports.periods import PERIOD_LTM, PERIOD_MONTH, PERIOD_YEAR, PERIOD_YTD, PeriodSpec, month_key
from apps.modules.reports.units import UNITS

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

UNIT_FORMATS = {
    "sum": '#,##0;[Red]-#,##0;"—"',
    "k": '#,##0,;[Red]-#,##0,;"—"',
    "m": '#,##0.0,,;[Red]-#,##0.0,,;"—"',
}
UNIT_LABELS = {"sum": "сум", "k": "тыс. сум", "m": "млн сум"}
RATIO_FORMAT = "0.0%"
PCT_DELTA_FORMAT = "+0.0%;-0.0%;0.0%"
PP_DELTA_FORMAT = '+0.0" п.п.";-0.0" п.п.";0.0" п.п."'
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"

REPORT_TITLES = {"pnl": "Отчёт о прибылях и убытках", "cashflow": "Отчёт о движении денежных средств"}
FILE_PREFIXES = {"pnl": "PnL", "cashflow": "Cashflow"}
SOURCE_LABELS = {
    "bank": "Банк",
    "cash": "Касса",
    "request": "Заявка",
    "invest_return": "Инвест. выплата",
    "unknown": "Операция",
}
OPERATION_HEADERS = (
    "Дата", "Раздел", "Статья", "Источник", "№ заявки", "Контрагент", "Описание", "Сумма, сум", "Месяц", "Амортизация",
)
OPERATION_WIDTHS = (12, 24, 28, 16, 10, 28, 48, 16, 10, 13)

GOOD = "107A55"
BAD = "C73A2E"
MUTED = "8A94A3"
HEADER_FILL = PatternFill("solid", fgColor="F2F5F9")
RESULT_FILL = PatternFill("solid", fgColor="EAF1FB")
THIN = Side(style="thin", color="9AA5B5")
DOUBLE = Side(style="double", color="9AA5B5")

# openpyxl keeps the whole workbook in memory; a longer list asks for a narrower period instead.
MAX_OPERATION_ROWS = 50_000

REPORT_HEADER_ROW = 8
LINES_HEADER_ROW = 8


@dataclass(frozen=True)
class ExportFile:
    filename: str
    content: bytes
    content_type: str = XLSX_CONTENT_TYPE


@dataclass(frozen=True)
class StatementExport:
    """The statement exactly as the page gets it (`build_statement_for_tenant`), plus the period's operations."""

    statement: dict[str, Any]
    operations: list[dict[str, Any]]
    section_labels: dict[str, str]
    units: str
    author: str


@dataclass(frozen=True)
class LinesExport:
    """The operations behind one report cell or a filtered list (items shaped like `statement/lines/`)."""

    report: str
    title: str
    date_from: date
    date_to: date
    query: str
    items: list[dict[str, Any]]
    total: str
    section_labels: dict[str, str]
    company: str
    author: str
    generated_at: datetime


class StatementRenderer(ABC):
    """Turns a statement export into a file body. Another format (PDF, CSV) is another subclass."""

    @abstractmethod
    def render(self, export: StatementExport) -> bytes:
        raise NotImplementedError


class LinesRenderer(ABC):
    """Turns an operation list into a file body."""

    @abstractmethod
    def render(self, export: LinesExport) -> bytes:
        raise NotImplementedError


class StatementXlsxRenderer(StatementRenderer):
    """«Отчёт» (statement with Excel outline groups), «Операции» (filterable table), «Параметры» (methodology)."""

    def render(self, export: StatementExport) -> bytes:
        workbook = Workbook()
        workbook.properties.creator = export.author
        sheet = workbook.active
        sheet.title = "Отчёт"
        _write_statement_sheet(sheet, export)
        _write_operations(workbook.create_sheet("Операции"), export.operations, export.section_labels, header_row=1)
        _write_parameters(workbook.create_sheet("Параметры"), export.statement.get("methodology") or [])
        return _workbook_bytes(workbook)


class LinesXlsxRenderer(LinesRenderer):
    """One sheet «Операции»: a header saying what the list is, then a filterable table."""

    def render(self, export: LinesExport) -> bytes:
        workbook = Workbook()
        workbook.properties.creator = export.author
        sheet = workbook.active
        sheet.title = "Операции"
        _put_text(sheet["A1"], export.company)
        sheet["A1"].font = Font(bold=True, size=14)
        _put_text(sheet["A2"], f"{REPORT_TITLES[export.report]} · {export.title}")
        sheet["A2"].font = Font(bold=True, size=12)
        _put_text(sheet["A3"], f"Период: {_dmy(export.date_from)} – {_dmy(export.date_to)}")
        _put_text(sheet["A4"], f"Поиск: «{export.query}»" if export.query else "Поиск: —")
        sheet["A5"] = "Итого, сум"
        sheet["B5"] = Decimal(export.total)
        sheet["B5"].number_format = MONEY_FORMAT
        for cell in (sheet["A5"], sheet["B5"]):
            cell.font = Font(bold=True)
        _put_text(sheet["A6"], _stamp(export.generated_at, export.author))
        _write_operations(sheet, export.items, export.section_labels, header_row=LINES_HEADER_ROW)
        _landscape(sheet, title_row=LINES_HEADER_ROW)
        return _workbook_bytes(workbook)


def statement_filename(report: str, spec: PeriodSpec, main_to: date, subdomain: str, today: date) -> str:
    """`PnL_<период>_<subdomain>_<YYYY-MM-DD>.xlsx` (`Cashflow_…` for the cash-flow report)."""
    return f"{FILE_PREFIXES[report]}_{_period_slug(spec, main_to)}_{subdomain}_{today.isoformat()}.xlsx"


def lines_filename(report: str, date_from: date, date_to: date, subdomain: str, today: date) -> str:
    return (
        f"{FILE_PREFIXES[report]}_operations_{date_from.isoformat()}_{date_to.isoformat()}"
        f"_{subdomain}_{today.isoformat()}.xlsx"
    )


def _period_slug(spec: PeriodSpec, main_to: date) -> str:
    if spec.kind == PERIOD_MONTH and spec.month:
        return spec.month
    if spec.kind == PERIOD_YEAR and spec.year:
        return str(spec.year)
    if spec.kind == PERIOD_YTD:
        return f"{main_to.year}-YTD"
    if spec.kind == PERIOD_LTM:
        return f"LTM-{month_key(main_to)}"
    return month_key(main_to)


@dataclass(frozen=True)
class _SheetRow:
    variant: str  # header | total | line | result | ratio | balance | spacer
    row: dict[str, Any] | None
    level: int


def _sheet_rows(rows: list[dict[str, Any]]) -> list[_SheetRow]:
    """Every group expanded: a header row, the children one outline level deeper, then «Итого …» below them."""
    children: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row["parent"] is not None:
            children[row["parent"]].append(row)
    out: list[_SheetRow] = []

    def walk(row: dict[str, Any], level: int) -> None:
        if row["kind"] == "group":
            out.append(_SheetRow("header", row, level + 1))
            for child in children[row["id"]]:
                walk(child, level + 1)
            out.append(_SheetRow("total", row, level))
            return
        out.append(_SheetRow(row["kind"], row, level))

    for row in rows:
        if row["parent"] is not None:
            continue
        if row.get("separator_before"):
            out.append(_SheetRow("spacer", None, 0))
        walk(row, 0)
    return out


def _write_statement_sheet(sheet: Worksheet, export: StatementExport) -> None:
    statement = export.statement
    meta = statement["meta"]
    columns = statement["columns"]
    main = _main_column(columns)
    _put_text(sheet["A1"], meta.get("company"))
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = REPORT_TITLES[statement["report"]]
    sheet["A2"].font = Font(bold=True, size=12)
    dates = f" ({_dmy_iso(main['from'])} – {_dmy_iso(main['to'])})" if main and main["from"] and main["to"] else ""
    _put_text(sheet["A3"], f"Период: {meta['period_label']}{dates}")
    _put_text(sheet["A4"], f"Сравнение: {_comparison_text(columns)}")
    _put_text(sheet["A5"], f"Единицы: {UNIT_LABELS[export.units]}")
    _put_text(sheet["A6"], _stamp(datetime.fromisoformat(meta["generated_at"]), export.author))

    header = REPORT_HEADER_ROW
    sheet.cell(row=header, column=1, value="Статья")
    for position, column in enumerate(columns, start=2):
        text = f"{column['label']}\n{column['sublabel']}" if column["sublabel"] else column["label"]
        _put_text(sheet.cell(row=header, column=position), text)
    for cell in sheet[header]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.cell(row=header, column=1).alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[header].height = 32

    number = header
    for item in _sheet_rows(statement["rows"]):
        number += 1
        _write_statement_row(sheet, number, item, columns, UNIT_FORMATS[export.units])

    for offset, note in enumerate(_notes(columns, meta.get("today")), start=2):
        sheet.cell(row=number + offset, column=1, value=note).font = Font(italic=True, color=MUTED)

    sheet.column_dimensions["A"].width = 46
    for position, column in enumerate(columns, start=2):
        sheet.column_dimensions[get_column_letter(position)].width = 12 if column["kind"] == "delta" else 16
    sheet.freeze_panes = f"B{header + 1}"
    sheet.sheet_properties.outlinePr = Outline(summaryBelow=True, summaryRight=True)
    _landscape(sheet, title_row=header)


def _write_statement_row(
    sheet: Worksheet, number: int, item: _SheetRow, columns: list[dict[str, Any]], money_format: str
) -> None:
    if item.variant == "spacer" or item.row is None:
        return
    row = item.row
    label = row["label"]
    if item.variant == "total":
        label = f"Итого {_lower_first(label)}"
    elif row.get("hint"):
        label = f"{label} ({row['hint']})"
    indent = row["depth"] + (1 if row["kind"] == "ratio" else 0)
    label_cell = sheet.cell(row=number, column=1)
    _put_text(label_cell, label)
    label_cell.alignment = Alignment(indent=indent)
    colours: dict[int, str] = {}
    if item.variant != "header":
        for position, column in enumerate(columns, start=2):
            cell = sheet.cell(row=number, column=position)
            if column["kind"] == "delta":
                value, number_format, colour = _delta_cell(row, row["deltas"].get(column["key"]))
                cell.value = value
                if number_format:
                    cell.number_format = number_format
                if colour:
                    colours[position] = colour
                continue
            cell.value = _decimal(row["values"].get(column["key"]))
            cell.number_format = RATIO_FORMAT if row["kind"] == "ratio" else money_format
    _style_row(sheet, number, item, len(columns) + 1, colours)
    if item.level:
        sheet.row_dimensions[number].outline_level = item.level


def _style_row(sheet: Worksheet, number: int, item: _SheetRow, width: int, colours: dict[int, str]) -> None:
    row = item.row or {}
    strong = bool(row.get("strong"))
    emphasised = item.variant == "result" or (item.variant == "balance" and strong)
    bold = item.variant in ("header", "total") or emphasised
    italic = item.variant == "ratio"
    top = THIN if item.variant == "total" else None
    bottom = DOUBLE if strong else None
    for position in range(1, width + 1):
        cell = sheet.cell(row=number, column=position)
        cell.font = Font(bold=bold, italic=italic, color=colours.get(position, MUTED if italic else None))
        if emphasised:
            cell.fill = RESULT_FILL
        if top or bottom:
            cell.border = Border(top=top, bottom=bottom)


def _delta_cell(row: dict[str, Any], delta: dict[str, str] | None) -> tuple[Decimal | None, str | None, str | None]:
    """Value, number format and font colour of a Δ cell; green or red by what the change means for this line."""
    if not delta:
        return None, None, None
    if "pp" in delta:
        value, number_format = Decimal(delta["pp"]), PP_DELTA_FORMAT
    else:
        value, number_format = Decimal(delta["pct"]), PCT_DELTA_FORMAT
    favorable = _favorable(row["polarity"], value)
    return value, number_format, None if favorable is None else (GOOD if favorable else BAD)


def _favorable(polarity: str, change: Decimal) -> bool | None:
    if change == 0 or polarity == "neutral":
        return None
    return change < 0 if polarity == "expense" else change > 0


def _main_column(columns: list[dict[str, Any]]) -> dict[str, Any] | None:
    """The period total, or the reporting month in a monthly report (same rule as the page)."""
    return next((c for c in columns if c["key"] == "total"), None) or next(
        (c for c in columns if c["kind"] == "period"), None
    )


def _comparison_text(columns: list[dict[str, Any]]) -> str:
    compare = next((c for c in columns if c["kind"] == "compare"), None)
    if compare is not None and compare["from"] and compare["to"]:
        return f"с тем же периодом прошлого года ({_dmy_iso(compare['from'])} – {_dmy_iso(compare['to'])})"
    if any(c["kind"] == "delta" for c in columns):
        return "с предыдущим месяцем и тем же месяцем прошлого года"
    return "без сравнения"


def _notes(columns: list[dict[str, Any]], today: str | None) -> list[str]:
    notes = [
        f"* {c['label'].replace('*', '')}: период не закрыт, данные по {_dmy_iso(c['to'])}."
        for c in columns
        if c["kind"] == "period" and c["partial"] and c["to"] and c["to"] == today
    ]
    if any(c["before_start"] for c in columns if c["kind"] != "delta"):
        notes.append("Пустые ячейки — период до начала учёта, данных нет.")
    return notes


def _write_operations(
    sheet: Worksheet, items: list[dict[str, Any]], section_labels: dict[str, str], *, header_row: int
) -> None:
    """A filterable Excel table: one row per operation, the amount exact in сум."""
    for position, title in enumerate(OPERATION_HEADERS, start=1):
        sheet.cell(row=header_row, column=position, value=title).font = Font(bold=True)
    shown = items[:MAX_OPERATION_ROWS]
    for offset, item in enumerate(shown, start=1):
        number = header_row + offset
        amortization = item.get("amortization")
        values = (
            date.fromisoformat(item["date"]),
            section_labels.get(item["line_id"].split(".")[0], ""),
            item["line_label"],
            SOURCE_LABELS.get(item["source"], SOURCE_LABELS["unknown"]),
            item["request_id"],
            item["counterparty"] or None,
            item["title"],
            Decimal(item["amount"]),
            item["date"][:7],
            f"{amortization['index']} из {amortization['count']}" if amortization else None,
        )
        for position, value in enumerate(values, start=1):
            cell = sheet.cell(row=number, column=position)
            if isinstance(value, str):
                _put_text(cell, value)
            else:
                cell.value = value
        sheet.cell(row=number, column=1).number_format = DATE_FORMAT
        sheet.cell(row=number, column=8).number_format = MONEY_FORMAT
    if shown:
        last = get_column_letter(len(OPERATION_HEADERS))
        table = Table(displayName="Operations", ref=f"A{header_row}:{last}{header_row + len(shown)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        sheet.add_table(table)
    else:
        # Excel repairs a table without data rows, so an empty list is a plain header and a note.
        sheet.cell(row=header_row + 1, column=1, value="Операций за период нет").font = Font(italic=True, color=MUTED)
    if len(items) > len(shown):
        note = (
            f"В файле первые {_count(len(shown))} из {_count(len(items))} операций. "
            "Сузьте период, чтобы выгрузить все."
        )
        sheet.cell(row=header_row + len(shown) + 2, column=1, value=note).font = Font(italic=True, color=MUTED)
    sheet.freeze_panes = f"A{header_row + 1}"
    for position, width in enumerate(OPERATION_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width


def _write_parameters(sheet: Worksheet, methodology: list[dict[str, str]]) -> None:
    sheet["A1"] = "Параметр"
    sheet["B1"] = "Как считается"
    for cell in (sheet["A1"], sheet["B1"]):
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for number, rule in enumerate(methodology, start=2):
        label_cell = sheet.cell(row=number, column=1)
        _put_text(label_cell, rule["label"])
        label_cell.font = Font(bold=True)
        text_cell = sheet.cell(row=number, column=2)
        _put_text(text_cell, rule["text"])
        text_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 100


def _landscape(sheet: Worksheet, *, title_row: int) -> None:
    """Landscape, one page wide, the table header repeated on every page, page numbers in the footer."""
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"{title_row}:{title_row}"
    sheet.oddFooter.center.text = "Стр. &P из &N"


def _put_text(cell: Any, value: Any) -> None:
    """User text stays text: control characters (which openpyxl rejects) are dropped, and a leading `=` or an
    error-like value such as `#N/A` is stored as a string, never as a live formula or an error cell."""
    if value is None or value == "":
        cell.value = None
        return
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    cell.data_type = "s"


def _count(value: int) -> str:
    return f"{value:,}".replace(",", " ")


def _decimal(value: Any) -> Decimal | None:
    return None if value is None or value == "" else Decimal(str(value))


def _dmy(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def _dmy_iso(value: str) -> str:
    return _dmy(date.fromisoformat(value))


def _lower_first(text: str) -> str:
    return text[:1].lower() + text[1:]


def _stamp(generated_at: datetime, author: str) -> str:
    return f"Выгружено: {generated_at.strftime('%d.%m.%Y %H:%M')} · {author}"


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
