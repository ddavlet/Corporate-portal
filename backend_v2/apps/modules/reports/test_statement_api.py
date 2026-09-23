from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import override_settings
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from apps.modules.reports.models import TenantReportSettings
from apps.modules.reports.test_fixtures import TODAY, sample_payload
from apps.modules.reports.xlsx_export import XLSX_CONTENT_TYPE
from apps.tenants.models import Tenant, TenantMembership, TenantModuleConfig, TenantUserRole

User = get_user_model()
FETCH = "apps.modules.reports.services.fetch_n8n_report_payload"
STATEMENT = "/api/reports/statement/"
LINES = "/api/reports/statement/lines/"
TEMPLATES = "/api/reports/templates/"
EXPORT = "/api/reports/statement/export/"
LINES_EXPORT = "/api/reports/statement/lines/export/"


@override_settings(BASE_DOMAIN="example.com", ALLOWED_HOSTS=["*"], TIME_ZONE="Asia/Tashkent")
@patch("django.utils.timezone.localdate", return_value=TODAY)
class StatementApiTests(APITestCase):
    def setUp(self):
        cache.clear()
        self.tenant = Tenant.objects.create(name="Демо Трейд", subdomain="stmt", is_active=True)
        TenantModuleConfig.objects.create(tenant=self.tenant, module_key="reports", is_enabled=True)
        self.admin = self._member("stmt_admin", TenantUserRole.ROLE_ADMIN)
        self.director = self._member("stmt_director", TenantUserRole.ROLE_DIRECTOR)
        self.accountant = self._member("stmt_accountant", TenantUserRole.ROLE_ACCOUNTANT)

    def _member(self, username, role):
        user = User.objects.create_user(username=username, password="x")
        TenantMembership.objects.create(tenant=self.tenant, user=user, is_active=True)
        TenantUserRole.objects.create(tenant=self.tenant, user=user, role=role)
        return user

    def _auth(self, user):
        token = str(RefreshToken.for_user(user).access_token)
        return {"HTTP_HOST": "stmt.example.com", "HTTP_AUTHORIZATION": f"Bearer {token}"}

    # templates
    def test_templates_default_for_tenant_without_settings(self, _today):
        res = self.client.get(TEMPLATES, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.data["default"], "classic")
        self.assertEqual([t["key"] for t in res.data["allowed"]], ["classic", "professional"])
        self.assertEqual([t["key"] for t in res.data["available"]], ["classic", "professional"])

    def test_templates_patch_is_admin_only(self, _today):
        body = {"default_template": "professional", "allowed_templates": ["professional"]}
        self.assertEqual(self.client.patch(TEMPLATES, body, format="json", **self._auth(self.director)).status_code, 403)
        res = self.client.patch(TEMPLATES, body, format="json", **self._auth(self.admin))
        self.assertEqual(res.status_code, 200, res.content)
        row = TenantReportSettings.objects.get(tenant=self.tenant)
        self.assertEqual((row.default_template, row.allowed_templates), ("professional", ["professional"]))

    def test_templates_patch_rejects_invalid_settings(self, _today):
        bad_default = {"default_template": "professional", "allowed_templates": ["classic"]}
        unknown = {"default_template": "classic", "allowed_templates": ["classic", "ghost"]}
        self.assertEqual(self.client.patch(TEMPLATES, bad_default, format="json", **self._auth(self.admin)).status_code, 400)
        self.assertEqual(self.client.patch(TEMPLATES, unknown, format="json", **self._auth(self.admin)).status_code, 400)

    # statement
    def test_statement_returns_professional_pnl(self, _today):
        with patch(FETCH, return_value=sample_payload()):
            res = self.client.get(
                STATEMENT,
                {"template": "professional", "report": "pnl", "period": "ytd", "compare": "yoy"},
                **self._auth(self.director),
            )
        self.assertEqual(res.status_code, 200, res.content)
        rev = next(row for row in res.data["rows"] if row["id"] == "rev")
        self.assertEqual(rev["values"]["total"], "3500.00")
        self.assertEqual([c["key"] for c in res.data["columns"]][-3:], ["total", "compare", "delta"])
        self.assertEqual(res.data["template"], "professional")
        self.assertEqual(res.data["meta"]["company"], "Демо Трейд")
        self.assertEqual(res.data["meta"]["period_label"], "янв – 23 сен 2026")
        self.assertEqual(res.data["warnings"], [])
        self.assertTrue(any(rule["label"] == "Период отчёта" for rule in res.data["methodology"]))

    def test_statement_rejects_classic_template(self, _today):
        res = self.client.get(STATEMENT, {"template": "classic", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 400)

    def test_statement_forbidden_when_template_not_allowed(self, _today):
        TenantReportSettings.objects.create(tenant=self.tenant, default_template="classic", allowed_templates=["classic"])
        res = self.client.get(STATEMENT, {"template": "professional", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 403)

    def test_statement_requires_reports_role(self, _today):
        res = self.client.get(STATEMENT, {"template": "professional", "report": "pnl"}, **self._auth(self.accountant))
        self.assertEqual(res.status_code, 403)

    def test_statement_rejects_months_before_2000(self, _today):
        params = {"template": "professional", "report": "pnl", "period": "month", "month": "0001-01"}
        res = self.client.get(STATEMENT, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 400)
        self.assertIn("month", res.data)

    def test_statement_survives_a_malformed_start_month_from_the_source(self, _today):
        payload = sample_payload()
        payload["metadata"]["start_month"] = "2024-13"
        payload["report_settings"]["start_month"] = "2024-13"
        with patch(FETCH, return_value=payload):
            res = self.client.get(STATEMENT, {"template": "professional", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content[:300])
        self.assertEqual(res.data["meta"]["start_month"], "2025-08")  # earliest operation month instead

    def test_statement_names_an_unknown_template(self, _today):
        res = self.client.get(STATEMENT, {"template": "ghost", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data["detail"], "Шаблон «ghost» не найден.")

    def test_lines_reject_a_line_the_report_does_not_have(self, _today):
        auth = self._auth(self.director)
        for line in ("ebit", "opex.deadbeef"):
            params = {"template": "professional", "report": "pnl", "line": line, "from": "2026-08-01", "to": "2026-08-31"}
            with patch(FETCH, return_value=sample_payload()):
                self.assertEqual(self.client.get(LINES, params, **auth).status_code, 404, line)
                self.assertEqual(self.client.get(LINES_EXPORT, params, **auth).status_code, 404, line)

    def test_statement_validates_period_params(self, _today):
        auth = self._auth(self.director)
        year_without_year = {"template": "professional", "report": "pnl", "period": "year"}
        future_month = {"template": "professional", "report": "pnl", "period": "month", "month": "2026-10"}
        self.assertEqual(self.client.get(STATEMENT, year_without_year, **auth).status_code, 400)
        self.assertEqual(self.client.get(STATEMENT, future_month, **auth).status_code, 400)

    def test_statement_warnings_only_for_admin(self, _today):
        unassigned = [{"purpose": "Аренда техники", "count": 3, "amount": "42600000.00"}]
        params = {"template": "professional", "report": "pnl"}
        with patch(FETCH, return_value=sample_payload()), patch(
            "apps.modules.reports.pnl_builder.compute_unassigned_payment_purposes", return_value=unassigned
        ):
            admin_res = self.client.get(STATEMENT, params, **self._auth(self.admin))
            director_res = self.client.get(STATEMENT, params, **self._auth(self.director))
        self.assertEqual(
            admin_res.data["warnings"],
            [{"code": "unassigned_purposes", "count": 3, "amount": "42600000.00", "purposes": ["Аренда техники"]}],
        )
        self.assertEqual(director_res.data["warnings"], [])

    def test_statement_upstream_failure_is_503(self, _today):
        with patch(FETCH, side_effect=RuntimeError("No tenant_report_settings for tenant_id=1")):
            res = self.client.get(STATEMENT, {"template": "professional", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 503)
        self.assertIn("tenant_report_settings", res.data["detail"])

    # lines
    def test_lines_total_matches_statement_cell(self, _today):
        auth = self._auth(self.director)
        with patch(FETCH, return_value=sample_payload()):
            statement = self.client.get(STATEMENT, {"template": "professional", "report": "pnl"}, **auth)
            lines = self.client.get(
                LINES,
                {"template": "professional", "report": "pnl", "line": "rev", "from": "2026-08-01", "to": "2026-08-31"},
                **auth,
            )
        rev = next(row for row in statement.data["rows"] if row["id"] == "rev")
        self.assertEqual(lines.status_code, 200, lines.content)
        self.assertEqual(lines.data["total"], rev["values"]["2026-08"])
        self.assertEqual(lines.data["count"], 2)

    def test_lines_search_and_paging(self, _today):
        params = {
            "template": "professional", "report": "pnl", "line": "opex", "from": "2026-01-01",
            "to": "2026-09-23", "q": "выставка", "page_size": 2,
        }
        with patch(FETCH, return_value=sample_payload()):
            res = self.client.get(LINES, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual((res.data["count"], res.data["total"], len(res.data["items"])), (3, "300.00", 2))
        self.assertEqual(res.data["items"][0]["amortization"], {"index": 3, "count": 3})
        self.assertEqual(res.data["items"][0]["line_label"], "Маркетинг")

    def test_lines_rejects_reversed_range(self, _today):
        params = {"template": "professional", "report": "pnl", "line": "rev", "from": "2026-09-01", "to": "2026-08-01"}
        res = self.client.get(LINES, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 400)

    def test_lines_without_line_list_every_section(self, _today):
        params = {"template": "professional", "report": "pnl", "from": "2026-08-01", "to": "2026-08-31"}
        with patch(FETCH, return_value=sample_payload()):
            res = self.client.get(LINES, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual((res.data["line"], res.data["count"], res.data["total"]), ("", 6, "2270.00"))

    def test_lines_filter_by_source(self, _today):
        params = {"template": "professional", "report": "pnl", "from": "2026-08-01", "to": "2026-08-31", "source": "bank"}
        with patch(FETCH, return_value=sample_payload()):
            res = self.client.get(LINES, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual((res.data["count"], res.data["total"]), (1, "1500.00"))

    def test_lines_rejects_unknown_source(self, _today):
        params = {"template": "professional", "report": "pnl", "from": "2026-08-01", "to": "2026-08-31", "source": "ghost"}
        res = self.client.get(LINES, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 400)

    # exports
    def test_export_returns_the_statement_workbook(self, _today):
        params = {"template": "professional", "report": "pnl", "period": "ytd", "compare": "yoy", "units": "k"}
        with patch(FETCH, return_value=sample_payload()):
            res = self.client.get(EXPORT, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content[:300])
        self.assertEqual(res["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertEqual(res["Content-Disposition"], 'attachment; filename="PnL_2026-YTD_stmt_2026-09-23.xlsx"')
        workbook = load_workbook(BytesIO(res.content))
        self.assertEqual(workbook.sheetnames, ["Отчёт", "Операции", "Параметры"])
        report = workbook["Отчёт"]
        self.assertEqual((report["A1"].value, report["A5"].value), ("Демо Трейд", "Единицы: тыс. сум"))
        self.assertTrue(report["A6"].value.endswith("· stmt_director"))
        # Header + the 10 operations of 1 Jan – 23 Sep 2026 across revenue, opex, other and investor payouts.
        self.assertEqual(workbook["Операции"].max_row, 11)

    def test_export_validates_units(self, _today):
        res = self.client.get(EXPORT, {"template": "professional", "report": "pnl", "units": "bn"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 400)

    def test_export_respects_allowed_templates(self, _today):
        TenantReportSettings.objects.create(tenant=self.tenant, default_template="classic", allowed_templates=["classic"])
        res = self.client.get(EXPORT, {"template": "professional", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 403)

    def test_export_upstream_failure_is_503(self, _today):
        with patch(FETCH, side_effect=RuntimeError("No tenant_report_settings for tenant_id=1")):
            res = self.client.get(EXPORT, {"template": "professional", "report": "pnl"}, **self._auth(self.director))
        self.assertEqual(res.status_code, 503)

    def test_lines_export_returns_the_cell_operations(self, _today):
        params = {"template": "professional", "report": "pnl", "line": "rev", "from": "2026-08-01", "to": "2026-08-31"}
        with patch(FETCH, return_value=sample_payload()):
            res = self.client.get(LINES_EXPORT, params, **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content[:300])
        self.assertEqual(
            res["Content-Disposition"], 'attachment; filename="PnL_operations_2026-08-01_2026-08-31_stmt_2026-09-23.xlsx"'
        )
        sheet = load_workbook(BytesIO(res.content))["Операции"]
        self.assertEqual(sheet["A2"].value, "Отчёт о прибылях и убытках · Выручка")
        self.assertEqual(sheet["B5"].value, 1800)
        self.assertEqual(sheet.max_row, 10)  # table header on row 8 + August's bank and cash receipts

    def test_lines_export_logs_its_duration(self, _today):
        params = {"template": "professional", "report": "pnl", "line": "rev", "from": "2026-08-01", "to": "2026-08-31"}
        with patch(FETCH, return_value=sample_payload()), self.assertLogs("apps.modules.reports.services", level="INFO") as logs:
            self.client.get(LINES_EXPORT, params, **self._auth(self.director))
        self.assertTrue(any("kind=lines" in line and "ms=" in line for line in logs.output), logs.output)

    # legacy endpoints keep working
    def test_legacy_pnl_keeps_shape_and_adds_section(self, _today):
        with patch("apps.modules.reports.views.fetch_n8n_report_payload", return_value=sample_payload()):
            res = self.client.get("/api/reports/pnl/", **self._auth(self.director))
        self.assertEqual(res.status_code, 200, res.content)
        for key in ("metadata", "totals", "monthly", "rows", "revenue", "operational_expenses", "other_expenses"):
            self.assertIn(key, res.data)
        self.assertTrue(all("section" in row for row in res.data["rows"]))

    def test_legacy_pnl_runtime_error_is_503(self, _today):
        with patch("apps.modules.reports.views.fetch_n8n_report_payload", side_effect=RuntimeError("n8n token missing")):
            res = self.client.get("/api/reports/pnl/", **self._auth(self.director))
        self.assertEqual(res.status_code, 503)
