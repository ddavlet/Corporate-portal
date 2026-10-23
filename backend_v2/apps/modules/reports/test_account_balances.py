from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from apps.modules.bank_expenses.models import BankExpense, BankRevenue
from apps.modules.cashier.models import CashExpense, CashRevenue
from apps.modules.reports.account_balances import (
    MOVEMENT_SOURCES,
    WalletNotFound,
    build_account_balances,
    list_balance_accounts,
)
from apps.modules.reports.xlsx_export import XLSX_CONTENT_TYPE, _operations
from apps.modules.wallets.calendar_utils import TASHKENT
from apps.modules.wallets.resolution import get_or_create_bank_wallet, get_or_create_cash_wallet
from apps.tenants.models import Tenant, TenantMembership, TenantModuleConfig, TenantUserRole

User = get_user_model()
TODAY = date(2026, 9, 23)
BALANCES = "/api/reports/balances/"
ACCOUNTS = "/api/reports/balances/accounts/"
EXPORT = "/api/reports/balances/export/"
ALL_TYPES = list(MOVEMENT_SOURCES)
# The opening balance is as of 1 January of the current year: API tests pin «today» so 2026 rows keep their numbers.
NOW = datetime(2026, 9, 23, 12, 0, tzinfo=TASHKENT)


class _Movements:
    def _bank_in(self, wallet, day, amount, purpose="Поступление"):
        return BankRevenue.objects.create(
            tenant=wallet.tenant, created_by=self.user, doc_date=day, process_date=day, doc_no="1",
            account_name="ООО Альфа", inn="1", account_no="1", mfo="1", kredit_turnover=Decimal(amount),
            payment_purpose=purpose, wallet=wallet,
        )

    def _bank_out(self, wallet, day, amount, purpose="Оплата"):
        return BankExpense.objects.create(
            tenant=wallet.tenant, created_by=self.user, row_no=1, doc_date=day, process_date=day,
            expense_year=day.year, expense_month=day.month, expense_day=day.day, debit_turnover=Decimal(amount),
            payment_purpose=purpose, wallet=wallet,
        )

    def _cash_in(self, wallet, at, amount, confirmed=True):
        return CashRevenue.objects.create(
            tenant=wallet.tenant, created_by=self.user, revenue_at=at, total_sum=Decimal(amount), operation="Выручка",
            confirmed=confirmed, wallet=wallet,
        )

    def _cash_out(self, wallet, at, amount, confirmed=True):
        return CashExpense.objects.create(
            tenant=wallet.tenant, created_by=self.user, external_id="e", title="Расход", amount=Decimal(amount),
            expense_at=at, expense_year=at.year, expense_month=at.month, expense_day=at.day, confirmed=confirmed,
            wallet=wallet,
        )


class AccountBalancesServiceTests(_Movements, TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(name="Acme", subdomain="balances", is_active=True)
        self.user = User.objects.create_user(username="balances_user", password="x")
        self.bank = get_or_create_bank_wallet(tenant=self.tenant)
        self.bank.opening_balance = Decimal("1000")
        self.bank.save(update_fields=["opening_balance"])
        self.cash = get_or_create_cash_wallet(tenant=self.tenant, currency="UZS")

    def _build(self, wallet, date_from, date_to, wallet_types=ALL_TYPES):
        return build_account_balances(
            tenant_id=self.tenant.id,
            wallet_id=wallet.id,
            date_from=date_from,
            date_to=date_to,
            wallet_types=wallet_types,
            today=TODAY,
        )

    def test_period_opening_closing_and_daily_balances(self):
        self._bank_in(self.bank, date(2026, 2, 10), "400")
        self._bank_in(self.bank, date(2026, 3, 5), "500")
        self._bank_out(self.bank, date(2026, 3, 5), "200")
        self._bank_out(self.bank, date(2026, 3, 10), "100")

        report = self._build(self.bank, date(2026, 3, 1), date(2026, 3, 31))

        self.assertEqual(Decimal(report["opening"]), Decimal("1400"))
        self.assertEqual(Decimal(report["income"]), Decimal("500"))
        self.assertEqual(Decimal(report["expense"]), Decimal("300"))
        self.assertEqual(Decimal(report["closing"]), Decimal("1600"))
        self.assertEqual([day["date"] for day in report["days"]], ["2026-03-05", "2026-03-10"])
        first = report["days"][0]
        self.assertEqual((Decimal(first["opening"]), Decimal(first["closing"])), (Decimal("1400"), Decimal("1700")))
        # Within a bank day income goes first; every operation carries the balance after it.
        self.assertEqual(
            [Decimal(op["balance_after"]) for op in first["operations"]], [Decimal("1900"), Decimal("1700")]
        )
        self.assertEqual(report["days"][1]["opening"], first["closing"])

    def test_balance_before_opening_day_is_derived_backwards(self):
        """Opening balance is as of 1 January: December's balances come from subtracting what moved after them."""
        self._bank_in(self.bank, date(2025, 12, 20), "300")

        report = self._build(self.bank, date(2025, 12, 1), date(2025, 12, 31))

        self.assertEqual(Decimal(report["opening"]), Decimal("700"))
        self.assertEqual(Decimal(report["closing"]), Decimal("1000"))

    def test_unconfirmed_cash_operations_are_listed_but_not_counted(self):
        self._cash_in(self.cash, datetime(2026, 4, 2, 10, 0, tzinfo=TASHKENT), "50")
        self._cash_out(self.cash, datetime(2026, 4, 2, 11, 0, tzinfo=TASHKENT), "20", confirmed=False)

        report = self._build(self.cash, date(2026, 4, 1), date(2026, 4, 30))

        self.assertEqual(Decimal(report["closing"]), Decimal("50"))
        self.assertEqual(Decimal(report["expense"]), Decimal("0"))
        self.assertEqual(report["unconfirmed"]["count"], 1)
        self.assertEqual(Decimal(report["unconfirmed"]["expense"]), Decimal("20"))
        pending = report["days"][0]["operations"][1]
        self.assertFalse(pending["counted"])
        self.assertIsNone(pending["balance_after"])

    def test_cash_day_follows_tashkent_calendar(self):
        # 00:30 on 2 April in Tashkent is still 1 April in UTC: the operation belongs to 2 April.
        self._cash_in(self.cash, datetime(2026, 4, 2, 0, 30, tzinfo=TASHKENT), "10")

        self.assertEqual(self._build(self.cash, date(2026, 4, 1), date(2026, 4, 1))["days"], [])
        report = self._build(self.cash, date(2026, 4, 2), date(2026, 4, 2))
        self.assertEqual([day["date"] for day in report["days"]], ["2026-04-02"])

    def test_as_of_date_starts_at_first_operation(self):
        self._bank_in(self.bank, date(2025, 11, 3), "300")
        self._bank_out(self.bank, date(2026, 5, 1), "50")

        report = self._build(self.bank, None, date(2026, 5, 31))

        self.assertTrue(report["since_first_operation"])
        self.assertEqual(report["date_from"], "2025-11-03")
        self.assertEqual(Decimal(report["opening"]), Decimal("700"))
        self.assertEqual(Decimal(report["closing"]), Decimal("950"))

    def test_as_of_date_cash_starts_at_effective_day_of_revenue(self):
        """A revenue entered on 1 March for 20 March starts the history on 20 March, not on the day it was entered."""
        revenue = self._cash_in(self.cash, datetime(2026, 3, 20, 10, 0, tzinfo=TASHKENT), "70")
        CashRevenue.objects.filter(id=revenue.id).update(created_at=datetime(2026, 3, 1, 9, 0, tzinfo=TASHKENT))
        self._cash_out(self.cash, datetime(2026, 3, 25, 10, 0, tzinfo=TASHKENT), "20")

        report = self._build(self.cash, None, date(2026, 3, 31))

        self.assertEqual(report["date_from"], "2026-03-20")
        self.assertEqual(Decimal(report["closing"]), Decimal("50"))

    def test_cash_revenue_without_date_falls_on_entry_day(self):
        revenue = self._cash_in(self.cash, None, "40")
        CashRevenue.objects.filter(id=revenue.id).update(created_at=datetime(2026, 5, 4, 15, 0, tzinfo=TASHKENT))

        report = self._build(self.cash, date(2026, 5, 1), date(2026, 5, 31))

        self.assertEqual([day["date"] for day in report["days"]], ["2026-05-04"])
        self.assertEqual(Decimal(report["income"]), Decimal("40"))

    def test_wallet_type_the_user_may_not_see_is_not_found(self):
        with self.assertRaises(WalletNotFound):
            self._build(self.bank, None, TODAY, wallet_types=["cash"])

    def test_wallet_of_another_company_is_not_found(self):
        other = Tenant.objects.create(name="Beta", subdomain="balances-other", is_active=True)
        foreign = get_or_create_bank_wallet(tenant=other)
        with self.assertRaises(WalletNotFound):
            self._build(foreign, None, TODAY)

    def test_accounts_list_bank_first(self):
        accounts = list_balance_accounts(tenant_id=self.tenant.id, wallet_types=ALL_TYPES)
        self.assertEqual([a["wallet_type"] for a in accounts], ["bank", "cash"])
        self.assertEqual(accounts[0]["type_label"], "Банк")
        only_cash = list_balance_accounts(tenant_id=self.tenant.id, wallet_types=["cash"])
        self.assertEqual([a["wallet_type"] for a in only_cash], ["cash"])


@override_settings(BASE_DOMAIN="example.com", ALLOWED_HOSTS=["*"])
class AccountBalancesApiTests(_Movements, APITestCase):
    def setUp(self):
        clock = patch("apps.modules.reports.account_balances.now_tashkent", return_value=NOW)
        clock.start()
        self.addCleanup(clock.stop)
        self.tenant = Tenant.objects.create(name="Демо", subdomain="balapi", is_active=True)
        for key in ("reports", "bank", "cash", "corporate_card"):
            TenantModuleConfig.objects.create(tenant=self.tenant, module_key=key, is_enabled=True)
        self.user = self._member("balapi_director", TenantUserRole.ROLE_DIRECTOR)
        self.bank = get_or_create_bank_wallet(tenant=self.tenant)

    def _member(self, username, role):
        user = User.objects.create_user(username=username, password="x")
        TenantMembership.objects.create(tenant=self.tenant, user=user, is_active=True)
        TenantUserRole.objects.create(tenant=self.tenant, user=user, role=role)
        return user

    def _auth(self, user=None):
        token = str(RefreshToken.for_user(user or self.user).access_token)
        return {"HTTP_HOST": "balapi.example.com", "HTTP_AUTHORIZATION": f"Bearer {token}"}

    def test_accounts_endpoint(self):
        res = self.client.get(ACCOUNTS, **self._auth())
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual([a["wallet_id"] for a in res.data["accounts"]], [self.bank.id])

    def test_balances_endpoint(self):
        self._bank_in(self.bank, date(2026, 3, 5), "500")
        res = self.client.get(
            BALANCES, {"wallet": self.bank.id, "date_from": "2026-03-01", "date_to": "2026-03-31"}, **self._auth()
        )
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(Decimal(res.data["income"]), Decimal("500"))
        self.assertEqual(len(res.data["days"]), 1)

    def test_balances_rejects_reversed_period(self):
        res = self.client.get(
            BALANCES, {"wallet": self.bank.id, "date_from": "2026-04-01", "date_to": "2026-03-01"}, **self._auth()
        )
        self.assertEqual(res.status_code, 400)

    def test_rejects_dates_out_of_range(self):
        for params in ({"date_to": "9999-12-31"}, {"date_from": "1999-12-31", "date_to": "2026-03-01"}):
            res = self.client.get(BALANCES, {"wallet": self.bank.id, **params}, **self._auth())
            self.assertEqual(res.status_code, 400, params)

    def test_investor_sees_no_accounts_without_bank_access(self):
        """Reports are open to investors, Bank is not: its accounts and their operations stay hidden."""
        investor = self._member("balapi_investor", TenantUserRole.ROLE_INVESTOR)
        res = self.client.get(ACCOUNTS, **self._auth(investor))
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.data["accounts"], [])
        res = self.client.get(BALANCES, {"wallet": self.bank.id, "date_to": "2026-03-31"}, **self._auth(investor))
        self.assertEqual(res.status_code, 404)

    def test_disabled_bank_module_hides_bank_accounts(self):
        cash = get_or_create_cash_wallet(tenant=self.tenant, currency="UZS")
        TenantModuleConfig.objects.filter(tenant=self.tenant, module_key="bank").update(is_enabled=False)
        res = self.client.get(ACCOUNTS, **self._auth())
        self.assertEqual([a["wallet_id"] for a in res.data["accounts"]], [cash.id])
        res = self.client.get(EXPORT, {"wallet": self.bank.id, "date_to": "2026-03-31"}, **self._auth())
        self.assertEqual(res.status_code, 404)

    def test_foreign_wallet_is_404(self):
        other = Tenant.objects.create(name="Beta", subdomain="balapi-other", is_active=True)
        foreign = get_or_create_bank_wallet(tenant=other)
        res = self.client.get(BALANCES, {"wallet": foreign.id, "date_to": "2026-03-31"}, **self._auth())
        self.assertEqual(res.status_code, 404)

    def test_export_has_balances_and_operations_sheets(self):
        self._bank_in(self.bank, date(2026, 3, 5), "500")
        self._bank_out(self.bank, date(2026, 3, 6), "200")
        res = self.client.get(
            EXPORT, {"wallet": self.bank.id, "date_from": "2026-03-01", "date_to": "2026-03-31"}, **self._auth()
        )
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res["Content-Type"], XLSX_CONTENT_TYPE)
        workbook = load_workbook(BytesIO(res.content))
        self.assertEqual(workbook.sheetnames, ["Остатки", "Операции"])
        operations = workbook["Операции"]
        self.assertEqual(operations.max_row, 3)
        self.assertEqual(Decimal(str(operations["I2"].value)), Decimal("500"))
        self.assertEqual(Decimal(str(operations["I3"].value)), Decimal("300"))
        balances = workbook["Остатки"]
        self.assertEqual(Decimal(str(balances["B10"].value)), Decimal("300"))
        self.assertEqual(balances["A13"].value.date(), date(2026, 3, 5))

    def test_requires_reports_module(self):
        TenantModuleConfig.objects.filter(tenant=self.tenant, module_key="reports").update(is_enabled=False)
        res = self.client.get(ACCOUNTS, **self._auth())
        self.assertEqual(res.status_code, 403)


class OperationsLabelTests(SimpleTestCase):
    def test_russian_plural(self):
        cases = {1: "1 операция", 3: "3 операции", 5: "5 операций", 11: "11 операций", 21: "21 операция", 1012: "1 012 операций"}
        for value, label in cases.items():
            self.assertEqual(_operations(value), label)
